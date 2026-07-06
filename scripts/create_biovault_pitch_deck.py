from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "BioVault_Pitch_Deck.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Arial"
FOOTER = "BioVault  Educational prototype"

NAVY = RGBColor(4, 18, 31)
NAVY_2 = RGBColor(7, 31, 48)
NAVY_3 = RGBColor(10, 45, 66)
TEAL = RGBColor(0, 184, 169)
AQUA = RGBColor(102, 231, 218)
WHITE = RGBColor(245, 250, 252)
MUTED = RGBColor(168, 197, 210)
MUTED_2 = RGBColor(117, 151, 166)
GRAY_LINE = RGBColor(64, 99, 115)
RED = RGBColor(232, 88, 88)
AMBER = RGBColor(248, 178, 72)
GREEN = RGBColor(72, 213, 157)


def inch(value):
    return Inches(value)


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def set_line(shape, color=GRAY_LINE, width=1.0, transparency=0):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def set_text(paragraph, size, color=WHITE, bold=False, align=None):
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    if align is not None:
        paragraph.alignment = align


def add_background(slide):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, inch(0), inch(0), inch(SLIDE_W), inch(SLIDE_H)
    )
    set_fill(bg, NAVY)
    bg.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, inch(0), inch(0), inch(0.13), inch(SLIDE_H)
    )
    set_fill(accent, TEAL)
    accent.line.fill.background()

    top = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, inch(0.13), inch(0), inch(SLIDE_W - 0.13), inch(0.04)
    )
    set_fill(top, NAVY_3)
    top.line.fill.background()


def add_footer(slide):
    box = slide.shapes.add_textbox(inch(0.55), inch(7.08), inch(5.0), inch(0.22))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = FOOTER
    set_text(p, 9, MUTED_2)


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(inch(0.62), inch(0.45), inch(12.1), inch(0.65))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = title
    set_text(p, 35, WHITE, True)
    if subtitle:
        sub = slide.shapes.add_textbox(inch(0.65), inch(1.08), inch(11.5), inch(0.38))
        stf = sub.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.text = subtitle
        set_text(sp, 17, MUTED)


def add_textbox(slide, x, y, w, h, text, size=18, color=MUTED, bold=False, align=None):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    set_text(p, size, color, bold, align)
    return box


def add_point(slide, x, y, text, size=18, color=MUTED, accent=TEAL):
    mark = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inch(x), inch(y + 0.1), inch(0.08), inch(0.08))
    set_fill(mark, accent)
    mark.line.fill.background()
    return add_textbox(slide, x + 0.18, y, 5.9, 0.46, text, size, color)


def add_box(
    slide,
    x,
    y,
    w,
    h,
    title,
    subtitle=None,
    fill=NAVY_2,
    line=GRAY_LINE,
    title_size=17,
    subtitle_size=12,
    title_color=WHITE,
    subtitle_color=MUTED,
    align=PP_ALIGN.CENTER,
):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, inch(x), inch(y), inch(w), inch(h))
    set_fill(shape, fill)
    set_line(shape, line, 1.1)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = inch(0.10)
    tf.margin_right = inch(0.10)
    tf.margin_top = inch(0.05)
    tf.margin_bottom = inch(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title
    set_text(p, title_size, title_color, True, align)
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        set_text(p2, subtitle_size, subtitle_color, False, align)
        p2.space_before = Pt(4)
    return shape


def add_arrow(slide, x, y, w, h, color=TEAL, transparency=0):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, inch(x), inch(y), inch(w), inch(h))
    set_fill(arrow, color, transparency)
    arrow.line.fill.background()
    return arrow


def add_down_arrow(slide, x, y, w, h, color=TEAL, transparency=0):
    arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, inch(x), inch(y), inch(w), inch(h))
    set_fill(arrow, color, transparency)
    arrow.line.fill.background()
    return arrow


def add_section_label(slide, x, y, text, color=AQUA):
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(3.8), inch(0.28))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text.upper()
    set_text(p, 10, color, True)
    return box


def add_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_footer(slide)
    return slide


def build_title_slide(prs):
    slide = add_slide(prs)
    add_section_label(slide, 0.70, 0.65, "BasedAI hackathon")

    box = slide.shapes.add_textbox(inch(0.68), inch(1.40), inch(9.6), inch(1.25))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "BioVault"
    set_text(p, 54, WHITE, True)

    add_textbox(
        slide,
        0.72,
        2.55,
        10.8,
        0.62,
        "Lineage-secured memory for AI science agents",
        26,
        AQUA,
        True,
    )
    add_textbox(
        slide,
        0.74,
        3.35,
        8.7,
        0.85,
        "A deterministic permission gate for shared scientific artifact memory.",
        20,
        MUTED,
    )

    add_box(
        slide,
        8.95,
        1.36,
        3.2,
        0.82,
        "0 model tokens",
        "in the permission path",
        fill=NAVY_3,
        line=TEAL,
        title_size=20,
        subtitle_size=12,
    )
    add_box(
        slide,
        8.95,
        2.44,
        3.2,
        0.82,
        "Lineage enforced",
        "across transitive sources",
        fill=NAVY_3,
        line=TEAL,
        title_size=20,
        subtitle_size=12,
    )
    add_box(
        slide,
        8.95,
        3.52,
        3.2,
        0.82,
        "Audit on every read",
        "request_id, provenance, latency",
        fill=NAVY_3,
        line=TEAL,
        title_size=18,
        subtitle_size=12,
    )

    add_textbox(
        slide,
        0.74,
        5.72,
        8.4,
        0.36,
        "Hackathon: BasedAI Enterprise Memory Governance at Scale",
        15,
        MUTED,
    )
    add_textbox(slide, 0.74, 6.15, 5.8, 0.35, "Don Aborah  @pkaysantana", 15, WHITE, True)
    return "BioVault: Lineage-Secured Memory for AI Science Agents"


def build_problem_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "Siloed copies break scientific memory governance")
    add_textbox(
        slide,
        0.75,
        1.34,
        5.8,
        0.65,
        "AI science agents assemble derived artifacts from sensitive sources.",
        22,
        WHITE,
        True,
    )
    points = [
        "Phase II memos, CRO handoffs, regulatory summaries, and scientific reports inherit source obligations.",
        "When a source is revoked, every artifact derived from it must quarantine automatically.",
        "External CROs must not read clinical memos they were never granted.",
        "Copying files into silos creates drift and stale permissions.",
    ]
    for idx, point in enumerate(points):
        add_point(slide, 0.78, 2.25 + idx * 0.72, point, 17)

    add_box(slide, 7.25, 1.45, 2.1, 0.68, "Sensitive source", "SAR / toxicity / AE", fill=NAVY_3, line=TEAL)
    add_arrow(slide, 9.48, 1.61, 0.82, 0.34, TEAL)
    add_box(slide, 10.45, 1.45, 2.0, 0.68, "Copied memo", "permission drift", fill=NAVY_3, line=AMBER)

    add_box(slide, 7.25, 2.73, 2.1, 0.68, "Revoked source", "access changes", fill=NAVY_3, line=RED)
    add_arrow(slide, 9.48, 2.89, 0.82, 0.34, AMBER)
    add_box(slide, 10.45, 2.73, 2.0, 0.68, "Stale artifact", "still readable", fill=NAVY_3, line=AMBER)

    add_box(slide, 7.25, 4.02, 2.1, 0.68, "External CRO", "narrow grant", fill=NAVY_3, line=TEAL)
    add_arrow(slide, 9.48, 4.18, 0.82, 0.34, RED)
    add_box(slide, 10.45, 4.02, 2.0, 0.68, "Clinical memo", "should deny", fill=NAVY_3, line=RED)

    add_textbox(
        slide,
        7.25,
        5.42,
        5.1,
        0.60,
        "The failure mode is not generation. It is uncontrolled retrieval of derived memory.",
        19,
        AQUA,
        True,
        PP_ALIGN.CENTER,
    )
    return "Problem: AI science agents create sensitive derived artifacts"


def build_insight_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "Governance must happen before content reaches a model")

    add_textbox(
        slide,
        1.05,
        1.55,
        11.2,
        0.74,
        "Access is a retrieval-layer decision, not a prompt instruction.",
        29,
        WHITE,
        True,
        PP_ALIGN.CENTER,
    )

    add_box(slide, 1.00, 3.05, 2.25, 0.95, "Request", "principal + artifact + operation", fill=NAVY_3, line=TEAL)
    add_arrow(slide, 3.48, 3.32, 0.80, 0.35, TEAL)
    add_box(slide, 4.55, 2.92, 2.45, 1.20, "Permission check", "SQL grants + lineage", fill=NAVY_3, line=AQUA)
    add_arrow(slide, 7.22, 3.32, 0.80, 0.35, TEAL)
    add_box(slide, 8.28, 2.66, 2.20, 0.82, "ALLOW", "read/query may proceed", fill=NAVY_3, line=GREEN, title_color=GREEN)
    add_down_arrow(slide, 9.05, 3.66, 0.56, 0.68, RED)
    add_box(slide, 8.28, 4.55, 2.20, 0.82, "DENY", "no plaintext returned", fill=NAVY_3, line=RED, title_color=RED)

    add_box(
        slide,
        10.85,
        3.30,
        1.42,
        0.82,
        "0 tokens",
        "for access",
        fill=NAVY_2,
        line=TEAL,
        title_size=18,
        subtitle_size=11,
    )
    add_textbox(
        slide,
        1.15,
        5.72,
        10.9,
        0.50,
        "If the gate denies the read, the model never receives the sensitive text.",
        19,
        AQUA,
        True,
        PP_ALIGN.CENTER,
    )
    return "Insight: Governance must happen before any model sees content"


def build_solution_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "BioVault gates shared scientific artifact memory")
    add_textbox(
        slide,
        0.78,
        1.28,
        11.7,
        0.44,
        "BioVault is a permission gate and demo dashboard. It is not an AI agent.",
        20,
        WHITE,
        True,
        PP_ALIGN.CENTER,
    )

    columns = [
        ("Permission gate", "Bearer capability tokens bind principal, artifact, and operation."),
        ("Lineage graph", "Derived reads require grants on the target and every transitive source."),
        ("Audit trail", "Each decision records request_id, provenance, and latency."),
    ]
    for i, (title, body) in enumerate(columns):
        x = 0.85 + i * 4.15
        add_box(slide, x, 2.28, 3.35, 1.25, title, body, fill=NAVY_3, line=TEAL, title_size=20, subtitle_size=14)

    add_box(
        slide,
        1.55,
        4.50,
        10.25,
        0.92,
        "Deterministic access: SQL permission check, sub-millisecond local path, 0 model tokens",
        "The model can be open-weight, closed, or absent. The permission layer stays model-free.",
        fill=NAVY_2,
        line=AQUA,
        title_size=19,
        subtitle_size=14,
    )
    add_textbox(
        slide,
        1.10,
        6.05,
        11.15,
        0.38,
        "Result: shared artifact memory can be queried without copying sensitive files into uncontrolled silos.",
        18,
        MUTED,
        False,
        PP_ALIGN.CENTER,
    )
    return "Solution: Permission gate + lineage + audit"


def build_how_it_works_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "Every read is checked against token, grants, and lineage")

    steps = [
        ("1", "Capability token", "principal, artifact, operation"),
        ("2", "Target grant", "deterministic SQL check"),
        ("3", "Source lineage", "all transitive parents granted"),
        ("4", "Decision + audit", "allow or deny with provenance"),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.74 + i * 3.15
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x), inch(1.68), inch(0.46), inch(0.46))
        set_fill(circle, TEAL)
        circle.line.fill.background()
        tf = circle.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = num
        set_text(p, 15, NAVY, True, PP_ALIGN.CENTER)

        add_box(slide, x + 0.05, 2.25, 2.45, 1.10, title, body, fill=NAVY_3, line=TEAL, title_size=18, subtitle_size=13)
        if i < len(steps) - 1:
            add_arrow(slide, x + 2.63, 2.62, 0.60, 0.28, TEAL)

    add_box(
        slide,
        1.22,
        4.28,
        4.80,
        1.12,
        "Revocation propagation",
        "Revoke one source -> quarantine every descendant artifact.",
        fill=NAVY_2,
        line=RED,
        title_size=20,
        subtitle_size=15,
        title_color=WHITE,
    )
    add_box(
        slide,
        7.28,
        4.28,
        4.80,
        1.12,
        "Query-time gate",
        "Denied reads return no plaintext and still write an audit event.",
        fill=NAVY_2,
        line=AQUA,
        title_size=20,
        subtitle_size=15,
    )
    return "How it works: Capability tokens, deterministic access, lineage, revocation, audit"


def build_architecture_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "Deterministic enforcement sits between agent and memory")

    add_box(slide, 0.75, 2.00, 2.25, 1.05, "AI science agent", "read/query request", fill=NAVY_3, line=TEAL)
    add_arrow(slide, 3.20, 2.36, 0.72, 0.32, TEAL)
    add_box(slide, 4.15, 1.72, 2.75, 1.60, "BioVault gate", "capability token\nSQL grants\nlineage traversal", fill=NAVY_3, line=AQUA, title_size=20, subtitle_size=13)

    add_arrow(slide, 7.10, 1.82, 0.72, 0.32, GREEN)
    add_box(slide, 8.05, 1.45, 3.05, 0.92, "ALLOW", "plaintext may enter optional model context", fill=NAVY_3, line=GREEN, title_color=GREEN)

    add_arrow(slide, 7.10, 3.10, 0.72, 0.32, RED)
    add_box(slide, 8.05, 2.78, 3.05, 0.92, "DENY", "no plaintext, audit still logs", fill=NAVY_3, line=RED, title_color=RED)

    add_box(slide, 4.22, 4.64, 2.62, 0.82, "Artifact store", "content + metadata", fill=NAVY_2, line=GRAY_LINE)
    add_box(slide, 7.20, 4.64, 2.62, 0.82, "Lineage DB", "sources + descendants", fill=NAVY_2, line=GRAY_LINE)
    add_box(slide, 10.18, 4.64, 2.20, 0.82, "Audit log", "request_id + latency", fill=NAVY_2, line=GRAY_LINE)

    add_textbox(slide, 1.00, 5.95, 11.25, 0.45, "Permission path: deterministic SQL, 0 model tokens, model-free enforcement.", 19, AQUA, True, PP_ALIGN.CENTER)
    return "Architecture: Agent -> BioVault gate -> allow or deny"


def build_demo_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "BVK-14 shows access, lineage, and revoke behavior")
    add_textbox(
        slide,
        0.78,
        1.25,
        11.75,
        0.38,
        "Demo programme: public target biology + internal SAR + toxicity + adverse-event memo -> Phase II readiness memo",
        17,
        MUTED,
        False,
        PP_ALIGN.CENTER,
    )

    beats = [
        ("1", "Nora Singh reads Phase II memo", "ALLOW"),
        ("2", "Owen Brooks reads Phase II memo", "DENY, no plaintext"),
        ("3", "Inspect lineage", "4 parents -> derived memo"),
        ("4", "CEO revokes adverse_event_memo", "quarantine propagates"),
        ("5", "Read Phase II memo after revoke", "DENY: derived_from_revoked_source"),
        ("6", "Audit log proves the path", "request_id, principal, provenance, latency"),
        ("7", "Permission path remains deterministic", "SQL check, 0 model tokens"),
    ]
    for i, (num, action, result) in enumerate(beats):
        col = i % 2
        row = i // 2
        x = 0.90 + col * 6.05
        y = 2.00 + row * 0.93
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, inch(x), inch(y + 0.06), inch(0.36), inch(0.36))
        set_fill(circle, TEAL if i not in (1, 4) else RED)
        circle.line.fill.background()
        ctf = circle.text_frame
        ctf.clear()
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = ctf.paragraphs[0]
        p.text = num
        set_text(p, 12, NAVY if i not in (1, 4) else WHITE, True, PP_ALIGN.CENTER)

        add_textbox(slide, x + 0.50, y - 0.03, 4.92, 0.26, action, 15, WHITE, True)
        add_textbox(slide, x + 0.50, y + 0.30, 4.92, 0.28, result, 14, AQUA if i not in (1, 4) else RED, False)

    add_box(
        slide,
        7.02,
        5.76,
        5.20,
        0.55,
        "Judge-facing proof point: deny paths never reveal plaintext.",
        None,
        fill=NAVY_2,
        line=AQUA,
        title_size=16,
    )
    return "Live demo: BVK-14 storyline with allow, deny, lineage, revoke"


def build_lineage_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "One revoked source quarantines the derived memo")

    sources = [
        ("public_target_paper", "public biology", TEAL),
        ("internal_sar_table", "restricted SAR", TEAL),
        ("toxicity_report", "safety source", TEAL),
        ("adverse_event_memo", "revoked by CEO", RED),
    ]
    for i, (name, subtitle, line) in enumerate(sources):
        y = 1.56 + i * 0.92
        add_box(slide, 0.82, y, 3.30, 0.58, name, subtitle, fill=NAVY_3, line=line, title_size=15, subtitle_size=10)
        add_arrow(slide, 4.32, y + 0.18, 2.62, 0.20, RED if line == RED else TEAL, transparency=0 if line == RED else 15)

    memo = add_box(
        slide,
        7.20,
        2.20,
        4.30,
        1.72,
        "phase2_readiness_memo",
        "quarantined after adverse_event_memo revoke",
        fill=NAVY_3,
        line=RED,
        title_size=22,
        subtitle_size=15,
        title_color=WHITE,
        subtitle_color=RED,
    )
    memo.fill.transparency = 5

    add_box(
        slide,
        1.15,
        5.70,
        10.9,
        0.72,
        "Rule: derived reads require target grant plus grants on every transitive source.",
        "A single revoked parent is enough to deny the descendant.",
        fill=NAVY_2,
        line=AQUA,
        title_size=18,
        subtitle_size=14,
    )
    return "Lineage diagram: 4 sources -> Phase II memo, then quarantine after revoke"


def build_mapping_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "BioVault maps directly to the governance rubric")

    headers = ("BasedAI expectation", "BioVault evidence")
    rows = [
        ("Source ACL / revocation event", "Simulated ACL and revoke flow for BVK-14 sources."),
        ("Retrieval-layer enforcement", "SQL gate runs before read/query; deny returns no plaintext."),
        ("No LLM permission path", "0 model tokens; permission layer is model-free."),
        ("Structured audit logs", "request_id, principal, provenance, decision, latency."),
        ("Performance evidence", "P99 < 200ms local permission benchmark."),
        ("Derived memory lineage", "Transitive source grants required for derived artifacts."),
        ("Revocation propagation", "Revoked source quarantines descendants automatically."),
        ("Open-weight compatible bonuses", "Temporal grant expiry and query-time gate support."),
    ]

    x0, y0 = 0.75, 1.35
    col_w = [4.25, 7.45]
    row_h = 0.55
    for c, header in enumerate(headers):
        add_box(
            slide,
            x0 + sum(col_w[:c]),
            y0,
            col_w[c],
            row_h,
            header,
            None,
            fill=TEAL,
            line=TEAL,
            title_size=16,
            title_color=NAVY,
        )
    for r, (left, right) in enumerate(rows):
        y = y0 + row_h + r * row_h
        fill = NAVY_3 if r % 2 == 0 else NAVY_2
        add_box(slide, x0, y, col_w[0], row_h, left, None, fill=fill, line=GRAY_LINE, title_size=13, title_color=WHITE, align=PP_ALIGN.LEFT)
        add_box(slide, x0 + col_w[0], y, col_w[1], row_h, right, None, fill=fill, line=GRAY_LINE, title_size=13, title_color=MUTED, align=PP_ALIGN.LEFT)

    add_textbox(slide, 0.88, 6.25, 11.45, 0.35, "The prototype demonstrates the governance path without claiming production IAM sync or production security.", 15, AQUA, True, PP_ALIGN.CENTER)
    return "BasedAI mapping: Requirements to BioVault evidence"


def build_why_matters_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "Scientific memory needs reproducible access boundaries")

    outcomes = [
        ("Safer collaboration", "External CROs see only artifacts and sources they are granted."),
        ("Auditable decisions", "Every allow and deny is reconstructable from request_id and provenance."),
        ("Model-agnostic control", "Open-weight agents can use the same gate because enforcement is outside the model."),
    ]
    for i, (title, body) in enumerate(outcomes):
        x = 0.92 + i * 4.05
        add_box(slide, x, 2.05, 3.25, 1.52, title, body, fill=NAVY_3, line=TEAL, title_size=20, subtitle_size=14)

    add_box(
        slide,
        1.38,
        4.74,
        10.55,
        0.95,
        "The product boundary is clear: BioVault governs memory access before science content reaches an agent.",
        "That makes revocation, lineage, and audit independent of prompt behavior.",
        fill=NAVY_2,
        line=AQUA,
        title_size=19,
        subtitle_size=14,
    )
    return "Why it matters: Safer AI science workflows and auditable memory governance"


def build_scope_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "This prototype is intentionally scoped")
    add_textbox(
        slide,
        0.85,
        1.25,
        11.65,
        0.42,
        "BioVault demonstrates the governance mechanism. It does not claim production security or clinical readiness.",
        19,
        WHITE,
        True,
        PP_ALIGN.CENTER,
    )

    limitations = [
        "Educational prototype only",
        "Not production security",
        "Not clinical decision support",
        "Local and ephemeral SQLite",
        "Simulated ACL and revocation event, not full external connector sync",
        "Not full Hirebase integration",
        "No production KMS/HSM yet",
        "No tamper-evident audit ledger yet",
    ]
    for i, item in enumerate(limitations):
        col = i % 2
        row = i // 2
        x = 1.00 + col * 5.85
        y = 2.05 + row * 0.78
        add_point(slide, x, y, item, 17, MUTED, accent=AMBER if i < 4 else TEAL)

    add_box(
        slide,
        1.35,
        5.78,
        10.65,
        0.55,
        "Next production steps: real connector sync, hardened key management, durable audit integrity, and deployment controls.",
        None,
        fill=NAVY_2,
        line=GRAY_LINE,
        title_size=16,
    )
    return "Honest scope: Prototype limitations and non-claims"


def build_close_slide(prs):
    slide = add_slide(prs)
    add_title(slide, "Govern scientific memory before the model reads it")

    add_textbox(
        slide,
        1.15,
        1.55,
        11.10,
        0.60,
        "BioVault is lineage-secured memory for AI science agents.",
        26,
        AQUA,
        True,
        PP_ALIGN.CENTER,
    )
    add_box(
        slide,
        1.20,
        2.75,
        3.35,
        1.10,
        "Review",
        "the deterministic permission path",
        fill=NAVY_3,
        line=TEAL,
        title_size=21,
        subtitle_size=14,
    )
    add_box(
        slide,
        5.00,
        2.75,
        3.35,
        1.10,
        "Run",
        "the BVK-14 revoke scenario",
        fill=NAVY_3,
        line=TEAL,
        title_size=21,
        subtitle_size=14,
    )
    add_box(
        slide,
        8.80,
        2.75,
        3.35,
        1.10,
        "Evaluate",
        "model-free memory governance",
        fill=NAVY_3,
        line=TEAL,
        title_size=21,
        subtitle_size=14,
    )

    add_textbox(slide, 2.55, 4.88, 8.3, 0.35, "Repo: github.com/pkaysantana/BioVault", 18, WHITE, True, PP_ALIGN.CENTER)
    add_textbox(slide, 2.55, 5.34, 8.3, 0.35, "Demo: [local/deploy URL TBD]", 18, MUTED, False, PP_ALIGN.CENTER)
    add_textbox(slide, 2.55, 6.05, 8.3, 0.35, "Don Aborah  @pkaysantana", 16, AQUA, True, PP_ALIGN.CENTER)
    return "Close / Ask: BioVault governs scientific memory before the model reads it"


def build_deck():
    prs = Presentation()
    prs.slide_width = inch(SLIDE_W)
    prs.slide_height = inch(SLIDE_H)
    prs.core_properties.title = "BioVault Pitch Deck"
    prs.core_properties.subject = "Lineage-secured memory for AI science agents"
    prs.core_properties.author = "Don Aborah"

    builders = [
        build_title_slide,
        build_problem_slide,
        build_insight_slide,
        build_solution_slide,
        build_how_it_works_slide,
        build_architecture_slide,
        build_demo_slide,
        build_lineage_slide,
        build_mapping_slide,
        build_why_matters_slide,
        build_scope_slide,
        build_close_slide,
    ]
    titles = [builder(prs) for builder in builders]
    prs.save(OUTPUT)
    return titles


if __name__ == "__main__":
    slide_titles = build_deck()
    print(f"Saved to: {OUTPUT}")
    print("Slides:")
    for idx, title in enumerate(slide_titles, start=1):
        print(f"{idx}. {title}")
